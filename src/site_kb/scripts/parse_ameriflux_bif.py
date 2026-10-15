"""Stage lossless BADM/BIF groups for review, without asserting site variables.

Grouping follows ecosim-agent's ameriflux-bif-timeseries-sorter; see the research
note for the audited revision and differences. No EcoSIM defaults are applied.
"""

import csv
from datetime import date, datetime, timezone
import hashlib
import json
from pathlib import Path
from typing import Optional

import typer

from site_kb.scripts.fetch_ameriflux import SITE_CODE, atomic_write

REQUIRED = ("SITE_ID", "GROUP_ID", "VARIABLE_GROUP", "VARIABLE", "DATAVALUE")
# Deliberately limited: unfamiliar groups require review, not a date-based guess.
OBSERVATIONS = {"GRP_AG_BIOMASS_CROP", "GRP_AG_PROD_CROP", "GRP_LAI", "GRP_HEIGHTC",
                "GRP_LMA", "GRP_BIOMASS_CHEM", "GRP_SOIL_CHEM", "GRP_SPP_O"}
METADATA = {"GRP_LOCATION", "GRP_CLIM_AVG", "GRP_IGBP", "GRP_DOI", "GRP_TEAM_MEMBER"}
app = typer.Typer()


def cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value)


def read_rows(path: Path, sheet: str = "AMF-BIF") -> tuple[list[dict], str | None]:
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as handle:
            matrix = list(csv.reader(handle))
        selected_sheet = None
    elif path.suffix.lower() in {".xlsx", ".xlsm"}:
        import openpyxl
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
        try:
            if sheet not in workbook.sheetnames:
                raise ValueError(f"Missing BIF sheet {sheet!r}; specify --sheet explicitly")
            matrix = []
            for cells in workbook[sheet].iter_rows():
                if any(cell.data_type == "f" for cell in cells):
                    raise ValueError("BIF input contains formulas; export source values before ingestion")
                matrix.append([cell_text(cell.value) for cell in cells])
            selected_sheet = sheet
        finally:
            workbook.close()
    else:
        raise ValueError("BIF input must be CSV, XLSX, or XLSM")
    if not matrix:
        raise ValueError("Empty BIF input")
    headers = matrix[0]
    if len(set(headers)) != len(headers) or not all(headers) or set(REQUIRED) - set(headers):
        raise ValueError(f"BIF headers must be unique and include {', '.join(REQUIRED)}")
    rows = []
    for row_number, row in enumerate(matrix[1:], start=2):
        if not any(row):
            continue
        if len(row) != len(headers):
            raise ValueError(f"BIF row {row_number}: column count differs from header")
        values = dict(zip(headers, row))
        if any(not values[k].strip() for k in REQUIRED if k != "DATAVALUE"):
            raise ValueError(f"BIF row {row_number}: empty record identity or variable")
        if not SITE_CODE.fullmatch(values["SITE_ID"]):
            raise ValueError(f"BIF row {row_number}: invalid SITE_ID")
        rows.append({"row_number": row_number, "values": values})
    if not rows:
        raise ValueError("BIF contains no data rows")
    return rows, selected_sheet


def normalize_date(raw: str) -> tuple[str | None, str | None]:
    """Keep year/month precision; never invent a day for an incomplete date."""
    value = raw.strip()
    if value in {"", "-9999", "NA"}:
        return None, None
    formats = [("%Y", 4, "year"), ("%Y%m", 6, "month"),
               ("%Y%m%d", 8, "day"), ("%Y%m%d%H%M", 12, "minute"),
               ("%Y%m%d%H%M%S", 14, "second")]
    for fmt, length, precision in formats:
        if value.isdigit() and len(value) == length:
            try:
                parsed = datetime.strptime(value, fmt)
            except ValueError:
                return None, None
            if parsed.strftime(fmt) != value:
                return None, None
            return {"year": value, "month": f"{value[:4]}-{value[4:]}",
                    "day": parsed.date().isoformat(),
                    "minute": parsed.isoformat(timespec="minutes"),
                    "second": parsed.isoformat(timespec="seconds")}[precision], precision
    try:
        if len(value) == 10:
            return date.fromisoformat(value).isoformat(), "day"
        if "T" in value:
            return datetime.fromisoformat(value).isoformat(), "datetime"
    except ValueError:
        pass
    return None, None


def reconstruct_records(rows: list[dict]) -> list[dict]:
    groups = {}
    for row in rows:
        values = row["values"]
        key = tuple(values[k] for k in ("SITE_ID", "VARIABLE_GROUP", "GROUP_ID"))
        groups.setdefault(key, []).append(row)
    records = []
    for (site_id, group, group_id), source_rows in sorted(groups.items()):
        kind = ("OBSERVATION" if group in OBSERVATIONS else
                "MANAGEMENT_EVENT" if group.startswith("GRP_DM_") else
                "METADATA" if group in METADATA else "UNCLASSIFIED")
        fields, dates, issues = {}, [], []
        for row in source_rows:
            variable, raw = (row["values"][k] for k in ("VARIABLE", "DATAVALUE"))
            fields.setdefault(variable, []).append(raw)
            if variable.endswith(("_DATE", "_DATE_START", "_DATE_END")):
                normalized, precision = normalize_date(raw)
                dates.append({"variable": variable, "raw_value": raw, "normalized": normalized,
                              "precision": precision, "source_row": row["row_number"]})
                if normalized is None and raw.strip() not in {"", "-9999", "NA"}:
                    issues.append(f"Unparsed date at row {row['row_number']}: {variable}")
        for variable, values in fields.items():
            if len(values) > 1:
                issues.append(f"Repeated variable retained without collapsing: {variable}")
        records.append({"site_id": site_id, "variable_group": group, "group_id": group_id,
                        "record_kind": kind,
                        "date_status": "DATED" if any(d["normalized"] for d in dates) else
                                       "UNPARSEABLE" if any(d["raw_value"].strip() not in {"", "-9999", "NA"} for d in dates) else "UNDATED",
                        "fields": fields, "dates": dates, "source_rows": source_rows,
                        "review_issues": issues})
    return records


def stage_bif(path: Path, raw_dir: Path, output_dir: Path, sheet: str = "AMF-BIF",
              source_url: str | None = None) -> Path:
    """Archive original bytes and write review JSON, outside the site assembly."""
    rows, selected_sheet = read_rows(path, sheet)
    payload = path.read_bytes()
    digest = hashlib.sha256(payload).hexdigest()
    archive = raw_dir / f"{digest}{path.suffix.lower()}"
    result = {"source_file": path.name, "source_sha256": digest,
              "archived_at": datetime.now(timezone.utc).isoformat(), "sheet": selected_sheet,
              "records": reconstruct_records(rows)}
    if source_url:
        result["source_url"] = source_url
    archive.parent.mkdir(parents=True, exist_ok=True)
    if archive.exists() and archive.read_bytes() != payload:
        raise ValueError(f"Existing BIF archive checksum mismatch: {archive}")
    if not archive.exists():
        archive.write_bytes(payload)
    # Include sheet identity so staging two sheets cannot overwrite each other.
    sheet_key = hashlib.sha256((selected_sheet or "CSV").encode()).hexdigest()[:12]
    output = output_dir / f"{digest}-{sheet_key}.json"
    atomic_write(output, json.dumps(result, ensure_ascii=False, indent=2) + "\n")
    return output


@app.command()
def main(
    path: Path = typer.Argument(..., exists=True, dir_okay=False),
    raw_dir: Path = typer.Option(Path("db/raw/ameriflux/bif")),
    output_dir: Path = typer.Option(Path("db/staged/ameriflux/bif")),
    sheet: str = typer.Option("AMF-BIF"),
    source_url: Optional[str] = typer.Option(None, help="Source download/record URL, if known"),
):
    """Stage a locally supplied BIF file; no downloads or site declarations."""
    output = stage_bif(path, raw_dir, output_dir, sheet, source_url)
    typer.echo(f"Staged BIF groups for review: {output}")


if __name__ == "__main__":
    app()
